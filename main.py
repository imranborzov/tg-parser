import csv
import io
import logging
from pathlib import Path
from fastapi import FastAPI, Request, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
import uvicorn
from contextlib import asynccontextmanager

BASE_DIR = Path(__file__).parent

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)

from src.database_manager import (
    init_db,
    list_instances,
    get_instance,
    create_instance,
    update_instance,
    delete_instance,
    count_instances,
    get_events,
    enqueue_joins,
    get_join_queue,
    get_join_queue_summary,
    clear_finished_joins,
)
from src.telegram_client import (
    start_all_clients_bg,
    stop_all_clients,
    reset_client,
    delete_client,
    send_code,
    verify_code,
    verify_2fa,
    is_authorized,
    send_to_webhook,
    pause_instance,
    resume_instance,
    is_instance_paused,
    normalize_link,
    pause_joins,
    resume_joins,
    is_joins_paused,
)
from typing import Optional

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Initialize SQLite database
    await init_db()
    # Connect every instance with credentials; start listening for authorized ones
    await start_all_clients_bg()
    try:
        yield
    finally:
        # Cleanup
        await stop_all_clients()

app = FastAPI(lifespan=lifespan)

# Mount static files (CSS/JS)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "src" / "static")), name="static")

# Templates
templates = Jinja2Templates(directory=str(BASE_DIR / "src" / "templates"))

class SetupData(BaseModel):
    channels: list[str]
    keywords: list[str]
    webhook_url: str
    tg_bot_token: str = ""
    tg_chat_id: str = ""
    match_cooldown: int = 60
    excluded_keywords: list[str] = []
    api_id: Optional[str] = None
    api_hash: Optional[str] = None
    phone: Optional[str] = None
    include_channels: bool = False

class InstanceCreate(BaseModel):
    name: str = "Untitled"

class InstanceRename(BaseModel):
    name: str

class AuthData(BaseModel):
    code: str
    phone_code_hash: str

class TwoFAData(BaseModel):
    password: str


@app.get("/", response_class=HTMLResponse)
async def read_root(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

# --- Instances ---

@app.get("/api/instances")
async def api_list_instances():
    return await list_instances()

@app.post("/api/instances")
async def api_create_instance(data: InstanceCreate):
    return await create_instance(data.name)

@app.get("/api/instances/{instance_id}")
async def api_get_instance(instance_id: int):
    inst = await get_instance(instance_id)
    if inst is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    return inst

@app.post("/api/instances/{instance_id}")
async def api_update_instance(instance_id: int, data: SetupData):
    existing = await get_instance(instance_id)
    if existing is None:
        raise HTTPException(status_code=404, detail="Instance not found")

    # Detect credential changes so we can rebuild the client only when needed
    # (rebuilding mid-listen would drop an active connection).
    creds_changed = (
        (data.api_id is not None and data.api_id.strip() != (existing.get("api_id") or "")) or
        (data.api_hash is not None and data.api_hash.strip() != (existing.get("api_hash") or "")) or
        (data.phone is not None and data.phone.strip() != (existing.get("phone") or ""))
    )

    await update_instance(
        instance_id,
        channels=data.channels,
        keywords=data.keywords,
        webhook_url=data.webhook_url,
        tg_bot_token=data.tg_bot_token,
        tg_chat_id=data.tg_chat_id,
        match_cooldown=data.match_cooldown,
        excluded_keywords=data.excluded_keywords,
        api_id=data.api_id,
        api_hash=data.api_hash,
        phone=data.phone,
        include_channels=data.include_channels,
    )
    if creds_changed:
        await reset_client(instance_id)
    return {"status": "success"}

@app.post("/api/instances/{instance_id}/rename")
async def api_rename_instance(instance_id: int, data: InstanceRename):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    await update_instance(instance_id, name=data.name)
    return {"status": "success"}

@app.delete("/api/instances/{instance_id}")
async def api_delete_instance(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    if await count_instances() <= 1:
        raise HTTPException(status_code=400, detail="Cannot delete the last instance.")
    await delete_client(instance_id)
    await delete_instance(instance_id)
    return {"status": "success"}

# --- Per-instance auth (each instance is its own Telegram account) ---

@app.get("/api/instances/{instance_id}/status")
async def api_instance_status(instance_id: int):
    inst = await get_instance(instance_id)
    if inst is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    has_credentials = bool((inst.get("api_id") or "").strip()
                           and (inst.get("api_hash") or "").strip()
                           and (inst.get("phone") or "").strip())
    authorized = await is_authorized(instance_id) if has_credentials else False
    return {"authorized": authorized, "has_credentials": has_credentials, "paused": is_instance_paused(instance_id)}

@app.post("/api/instances/{instance_id}/pause")
async def api_pause_instance(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    pause_instance(instance_id)
    return {"status": "success", "paused": True}

@app.post("/api/instances/{instance_id}/resume")
async def api_resume_instance(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    resume_instance(instance_id)
    return {"status": "success", "paused": False}

@app.post("/api/instances/{instance_id}/auth/send_code")
async def api_send_code(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    try:
        phone_code_hash = await send_code(instance_id)
        return {"status": "success", "phone_code_hash": phone_code_hash}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.post("/api/instances/{instance_id}/auth/verify_code")
async def api_verify_code(instance_id: int, data: AuthData):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    return await verify_code(instance_id, data.code, data.phone_code_hash)

@app.post("/api/instances/{instance_id}/auth/verify_2fa")
async def api_verify_2fa(instance_id: int, data: TwoFAData):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    return await verify_2fa(instance_id, data.password)

@app.get("/api/instances/{instance_id}/events")
async def api_instance_events(instance_id: int):
    return await get_events(instance_id)

@app.post("/api/instances/{instance_id}/webhook/test")
async def api_test_webhook(instance_id: int):
    inst = await get_instance(instance_id)
    if inst is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    webhook_url = inst.get("webhook_url", "")
    if not webhook_url:
        return {"status": "error", "message": "No webhook URL configured."}
    sample = {
        "channel_id": "-1001234567890",
        "channel_name": "Test Channel",
        "message_id": 1,
        "message_text": "This is a test message from TG Parser.",
        "message_link": "https://t.me/c/1234567890/1",
        "matched_keyword": "test",
        "date": "2024-01-01T00:00:00+00:00",
        "sender_id": "123456789",
        "sender_username": "example_user",
        "sender_name": "Example User",
        "sender_link": "https://t.me/example_user",
    }
    await send_to_webhook(webhook_url, sample)
    return {"status": "success"}

# --- Bulk channel upload (CSV / XLSX) ---

MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB


def _cells_from_csv(data: bytes):
    text = data.decode("utf-8-sig", errors="replace")
    for row in csv.reader(io.StringIO(text)):
        for cell in row:
            yield cell


def _cells_from_xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    yield str(cell)
    wb.close()


def _parse_links(filename: str, data: bytes) -> list[dict]:
    """Extract normalized join jobs from an uploaded file, deduped by identifier."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        cells = _cells_from_xlsx(data)
    elif name.endswith(".csv"):
        cells = _cells_from_csv(data)
    else:
        raise ValueError("Unsupported file type — upload a .csv or .xlsx file.")

    seen = set()
    jobs = []
    for cell in cells:
        job = normalize_link(cell)
        if job and job["identifier"] not in seen:
            seen.add(job["identifier"])
            jobs.append(job)
    return jobs


@app.post("/api/instances/{instance_id}/channels/upload")
async def api_upload_channels(instance_id: int, file: UploadFile = File(...)):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="File too large (max 5 MB).")
    try:
        jobs = _parse_links(file.filename, data)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read that file.")
    if not jobs:
        return {"status": "success", "found": 0, "queued": 0,
                "message": "No Telegram links found in the file."}
    queued = await enqueue_joins(instance_id, jobs)
    return {"status": "success", "found": len(jobs), "queued": queued,
            "skipped": len(jobs) - queued}


@app.get("/api/instances/{instance_id}/join_queue")
async def api_join_queue(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    return {
        "summary": await get_join_queue_summary(instance_id),
        "jobs": await get_join_queue(instance_id),
        "paused": is_joins_paused(instance_id),
    }


@app.post("/api/instances/{instance_id}/join_queue/clear")
async def api_clear_join_queue(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    removed = await clear_finished_joins(instance_id)
    return {"status": "success", "removed": removed}


@app.post("/api/instances/{instance_id}/join_queue/pause")
async def api_pause_joins(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    pause_joins(instance_id)
    return {"status": "success", "paused": True}


@app.post("/api/instances/{instance_id}/join_queue/resume")
async def api_resume_joins(instance_id: int):
    if await get_instance(instance_id) is None:
        raise HTTPException(status_code=404, detail="Instance not found")
    resume_joins(instance_id)
    return {"status": "success", "paused": False}

if __name__ == "__main__":
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
